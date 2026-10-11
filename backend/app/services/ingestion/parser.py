import io
import re
from abc import ABC, abstractmethod
import pymupdf
import docx
import chardet

from app.core.errors import (
    DocumentParsingError,
    EmptyDocumentError,
    UnsupportedEncodingError,
    UnsupportedFileTypeError,
)
from app.models.document import ParsedBlock, ParsedDocument, DocumentType


class BaseParser(ABC):
    """Abstract base class for document parsers."""

    @abstractmethod
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        """Parses raw document bytes and emits a normalized ParsedDocument."""
        pass


class PDFParser(BaseParser):
    """High-fidelity PDF parser using PyMuPDF (fitz) with physical page tracking."""

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        if not content or len(content.strip()) == 0:
            raise EmptyDocumentError(filename)

        try:
            doc = pymupdf.open(stream=content, filetype="pdf")
        except Exception as e:
            raise DocumentParsingError(filename, f"Invalid or corrupt PDF file: {e}")

        try:
            total_pages = doc.page_count
            if total_pages == 0:
                raise EmptyDocumentError(filename)

            blocks: list[ParsedBlock] = []
            total_chars = 0

            for page in doc:
                page_num = page.number + 1  # 1-indexed physical page
                raw_blocks = page.get_text("blocks")

                for b in raw_blocks:
                    # b: (x0, y0, x1, y1, text, block_no, block_type)
                    if len(b) >= 7 and b[6] == 0:  # block_type == 0 means text
                        text = b[4].strip()
                        if text:
                            clean_text = re.sub(r"[ \t]+", " ", text)
                            total_chars += len(clean_text)
                            blocks.append(
                                ParsedBlock(
                                    text=clean_text,
                                    page_number=page_num,
                                    metadata={
                                        "bbox": [round(b[0], 2), round(b[1], 2), round(b[2], 2), round(b[3], 2)],
                                        "block_no": b[5],
                                    },
                                )
                            )

            if not blocks or total_chars == 0:
                raise EmptyDocumentError(filename)

            return ParsedDocument(
                filename=filename,
                file_type=DocumentType.PDF.value,
                char_count=total_chars,
                blocks=blocks,
                page_count=total_pages,
                metadata={"title": doc.metadata.get("title") or filename},
            )
        finally:
            doc.close()


class DOCXParser(BaseParser):
    """DOCX parser using python-docx with section hierarchy and table support.
    Does NOT invent physical page numbers; uses exact section and paragraph/table tracking.
    """

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        if not content or len(content.strip()) == 0:
            raise EmptyDocumentError(filename)

        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise DocumentParsingError(filename, f"Invalid or corrupt DOCX file: {e}")

        blocks: list[ParsedBlock] = []
        total_chars = 0
        heading_levels: dict[int, str] = {}
        para_counter = 0
        table_counter = 0

        try:
            items = list(doc.iter_inner_content())
        except Exception as e:
            raise DocumentParsingError(filename, f"Failed to extract document contents: {e}")

        for item in items:
            if isinstance(item, docx.text.paragraph.Paragraph):
                text = item.text.strip()
                if not text:
                    continue

                style_name = (item.style.name or "").lower() if item.style else ""
                heading_match = re.search(r"heading\s*(\d)", style_name)

                if heading_match:
                    level = int(heading_match.group(1))
                    heading_levels[level] = text
                    # Clear deeper heading levels
                    heading_levels = {k: v for k, v in heading_levels.items() if k <= level}
                    current_heading = " > ".join(heading_levels[k] for k in sorted(heading_levels.keys()))
                else:
                    current_heading = " > ".join(heading_levels[k] for k in sorted(heading_levels.keys())) or "Document Body"

                para_counter += 1
                clean_text = re.sub(r"[ \t]+", " ", text)
                total_chars += len(clean_text)
                blocks.append(
                    ParsedBlock(
                        text=clean_text,
                        heading_path=current_heading,
                        paragraph_number=para_counter,
                        metadata={"style": item.style.name if item.style else "Normal"},
                    )
                )

            elif isinstance(item, docx.table.Table):
                table_counter += 1
                current_heading = " > ".join(heading_levels[k] for k in sorted(heading_levels.keys())) or "Document Body"

                # Convert table to markdown representation
                table_lines: list[str] = []
                for row_idx, row in enumerate(item.rows):
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_lines.append("| " + " | ".join(cells) + " |")
                    if row_idx == 0 and len(item.columns) > 0:
                        table_lines.append("| " + " | ".join(["---"] * len(item.columns)) + " |")

                table_text = "\n".join(table_lines).strip()
                if table_text:
                    total_chars += len(table_text)
                    blocks.append(
                        ParsedBlock(
                            text=table_text,
                            heading_path=current_heading,
                            table_number=table_counter,
                            metadata={"rows": len(item.rows), "cols": len(item.columns)},
                        )
                    )

        if not blocks or total_chars == 0:
            raise EmptyDocumentError(filename)

        return ParsedDocument(
            filename=filename,
            file_type=DocumentType.DOCX.value,
            char_count=total_chars,
            blocks=blocks,
            page_count=None,  # Not physically determinable in DOCX
            metadata={"paragraph_count": para_counter, "table_count": table_counter},
        )


class TXTParser(BaseParser):
    """UTF-8 text parser with chardet fallback and line number preservation."""

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        if not content or len(content.strip()) == 0:
            raise EmptyDocumentError(filename)

        # Attempt UTF-8 decoding first
        decoded_text: str | None = None
        try:
            decoded_text = content.decode("utf-8")
        except UnicodeDecodeError:
            detected = chardet.detect(content)
            encoding = detected.get("encoding")
            confidence = detected.get("confidence") or 0.0

            if encoding and confidence >= 0.7:
                try:
                    decoded_text = content.decode(encoding)
                except Exception:
                    pass

            if decoded_text is None:
                raise UnsupportedEncodingError(filename, encoding)

        # Normalize line breaks
        normalized_text = decoded_text.replace("\r\n", "\n").replace("\r", "\n")
        lines = normalized_text.split("\n")

        blocks: list[ParsedBlock] = []
        total_chars = 0
        current_block_lines: list[str] = []
        start_line = 1

        for line_idx, line in enumerate(lines, start=1):
            stripped = line.strip()
            if stripped:
                if not current_block_lines:
                    start_line = line_idx
                current_block_lines.append(stripped)
            else:
                if current_block_lines:
                    block_text = "\n".join(current_block_lines)
                    total_chars += len(block_text)
                    blocks.append(
                        ParsedBlock(
                            text=block_text,
                            line_start=start_line,
                            line_end=line_idx - 1,
                        )
                    )
                    current_block_lines = []

        # Flush final block
        if current_block_lines:
            block_text = "\n".join(current_block_lines)
            total_chars += len(block_text)
            blocks.append(
                ParsedBlock(
                    text=block_text,
                    line_start=start_line,
                    line_end=len(lines),
                )
            )

        if not blocks or total_chars == 0:
            raise EmptyDocumentError(filename)

        return ParsedDocument(
            filename=filename,
            file_type=DocumentType.TXT.value,
            char_count=total_chars,
            blocks=blocks,
            page_count=None,
            metadata={"line_count": len(lines)},
        )


class ParserFactory:
    """Factory creating appropriate parser based on file extension."""

    @staticmethod
    def get_parser(filename: str) -> BaseParser:
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext == DocumentType.PDF.value:
            return PDFParser()
        elif ext == DocumentType.DOCX.value:
            return DOCXParser()
        elif ext == DocumentType.TXT.value:
            return TXTParser()
        else:
            raise UnsupportedFileTypeError(ext, {DocumentType.PDF.value, DocumentType.DOCX.value, DocumentType.TXT.value})
