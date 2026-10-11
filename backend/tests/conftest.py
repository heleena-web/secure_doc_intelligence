import io
import sys
from pathlib import Path

backend_dir = Path(__file__).resolve().parent.parent
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

from typing import AsyncGenerator
import pytest
import pytest_asyncio
import docx
import pymupdf
from httpx import AsyncClient, ASGITransport

from app.core.config import settings
from app.main import app
from app.storage.database import init_db


from datetime import datetime, timezone
from app.core.security import create_access_token, hash_password
from app.models.auth import UserRecord, UserRole
from app.services.retrieval.search_service import search_service
from app.storage.database import get_db_connection
from app.storage.repository import DocumentRepository


@pytest.fixture(autouse=True)
def setup_test_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """Overrides storage, database, and vector paths to use isolated temporary directories."""
    test_db = tmp_path / "test_app.db"
    test_storage = tmp_path / "storage"
    test_uploads = tmp_path / "uploads"
    test_chroma = tmp_path / "chroma"
    test_storage.mkdir(parents=True, exist_ok=True)
    test_uploads.mkdir(parents=True, exist_ok=True)
    test_chroma.mkdir(parents=True, exist_ok=True)

    monkeypatch.setattr(settings, "SQLITE_DB_PATH", test_db)
    monkeypatch.setattr(settings, "STORAGE_DIR", test_storage)
    monkeypatch.setattr(settings, "UPLOAD_DIR", test_uploads)
    monkeypatch.setattr(settings, "CHROMA_PERSIST_DIR", test_chroma)

    search_service.clear()


@pytest_asyncio.fixture
async def unauthenticated_client() -> AsyncGenerator[AsyncClient, None]:
    """Provides an unauthenticated HTTP async client for FastAPI tests."""
    await init_db(settings.SQLITE_DB_PATH)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield client


@pytest_asyncio.fixture
async def async_client() -> AsyncGenerator[AsyncClient, None]:
    """Provides an authenticated HTTP async client with default admin credentials."""
    await init_db(settings.SQLITE_DB_PATH)
    test_user_id = "test_default_admin_id"
    async with get_db_connection(settings.SQLITE_DB_PATH) as db:
        repo = DocumentRepository(db)
        if not await repo.get_user_by_id(test_user_id):
            now = datetime.now(timezone.utc).isoformat()
            user = UserRecord(
                id=test_user_id,
                username="test_admin",
                email="admin@test.com",
                hashed_password=hash_password("adminpassword123"),
                full_name="Default Test Admin",
                is_active=True,
                is_superuser=True,
                created_at=now,
            )
            await repo.create_user(user)
            await repo.add_workspace_member("default", test_user_id, UserRole.OWNER)

    token = create_access_token(test_user_id, claims={"username": "test_admin"})
    transport = ASGITransport(app=app)
    headers = {"Authorization": f"Bearer {token}"}
    async with AsyncClient(transport=transport, base_url="http://test", headers=headers) as client:
        yield client


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """Creates a multi-page PDF with physical page content."""
    doc = pymupdf.open()
    # Page 1
    p1 = doc.new_page()
    p1.insert_text((50, 72), "Annual Security & Compliance Report - Page 1")
    p1.insert_text((50, 100), "All employees must complete mandatory security training annually.")
    # Page 2
    p2 = doc.new_page()
    p2.insert_text((50, 72), "Access Control Policies - Page 2")
    p2.insert_text((50, 100), "Zero-trust network architecture requires least-privilege token access.")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Creates a DOCX with hierarchical headings, paragraphs, and a table."""
    doc = docx.Document()
    doc.add_heading("Corporate Security Framework", level=1)
    doc.add_paragraph("This document outlines the core protocols governing data encryption and tenant isolation.")
    
    doc.add_heading("Section 2: Data Retention Guidelines", level=2)
    doc.add_paragraph("Audit logs must be retained for a minimum of 7 years in write-once-read-many storage.")
    
    # Add a table
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Classification"
    table.cell(0, 1).text = "Retention Period"
    table.cell(1, 0).text = "Public"
    table.cell(1, 1).text = "1 Year"
    table.cell(2, 0).text = "Restricted"
    table.cell(2, 1).text = "7 Years"
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_txt_bytes() -> bytes:
    """Creates a valid UTF-8 text document with multiple lines and paragraphs."""
    content = (
        "Server Configuration and Hardening Guide\n\n"
        "1. SSH configuration requires disabling root login and password authentication.\n"
        "2. All database connections must utilize TLS 1.3 encryption.\n\n"
        "3. Automated security patches must be scheduled weekly.\n"
    )
    return content.encode("utf-8")
